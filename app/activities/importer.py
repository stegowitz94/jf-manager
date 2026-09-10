from datetime import datetime, time
from io import BytesIO
from openpyxl import load_workbook
from django.utils import timezone

REQUIRED={"begin","ende","thema"}

def _dt(value):
    if isinstance(value, datetime): return value
    if not value: return None
    for fmt in ("%d.%m.%Y %H:%M","%Y-%m-%d %H:%M:%S","%Y-%m-%d %H:%M"):
        try: return datetime.strptime(str(value).strip(),fmt)
        except ValueError: pass
    raise ValueError(f"Ungültiges Datum: {value}")

def parse_divera(content):
    wb=load_workbook(BytesIO(content),data_only=True,read_only=True); ws=wb.active
    headers=[str(c.value or "").strip().lower() for c in next(ws.iter_rows())]
    missing=REQUIRED-set(headers)
    if missing: raise ValueError("Fehlende Spalten: "+", ".join(sorted(missing)))
    idx={h:i for i,h in enumerate(headers)}; rows=[]
    for number,row in enumerate(ws.iter_rows(values_only=True),start=2):
        if not any(row): continue
        try:
            start=_dt(row[idx["begin"]]); end=_dt(row[idx["ende"]])
            if timezone.is_naive(start): start=timezone.make_aware(start)
            if end and timezone.is_naive(end): end=timezone.make_aware(end)
            rows.append({"row":number,"starts_at":start.isoformat(),"ends_at":end.isoformat() if end else None,"title":str(row[idx["thema"]] or "").strip(),"description":str(row[idx.get("text",-1)] or "").strip() if "text" in idx else "","location":str(row[idx.get("ort",-1)] or "").strip() if "ort" in idx else ""})
        except Exception as exc: rows.append({"row":number,"error":str(exc)})
    return rows
