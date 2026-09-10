from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Guardian, Member


IMPORT_FIELDS = [
    ("sequential_number", "Laufende Nummer", False),
    ("first_name", "Vorname", True),
    ("last_name", "Nachname", True),
    ("gender", "Geschlecht", False),
    ("birth_date", "Geburtsdatum", True),
    ("entry_date", "Beitrittsdatum", False),
    ("jf_card_number", "JF-Ausweisnummer", False),
    ("activity_status", "Aktivitätsstatus", False),
    ("street", "Straße und Hausnummer", False),
    ("postal_code", "PLZ", False),
    ("city", "Ort", False),
    ("phone", "Telefon Mitglied", False),
    ("mobile", "Mobil Mitglied", False),
    ("email", "E-Mail Mitglied", False),
    ("photo_video_permission", "Foto-/Videoeinwilligung", False),
    ("divera_permission", "DIVERA-Einwilligung", False),
    ("school_or_employer", "Schule oder Arbeitgeber", False),
    ("vocational_training", "Berufsausbildung", False),
    ("career_goal", "Voraussichtliches Berufsziel", False),
    ("other_organizations", "Weitere Vereine oder Organisationen", False),
    ("health_notes", "Gesundheitliche Hinweise (sensibel)", False),
    ("swimming_status", "Schwimmfähigkeit (sensibel)", False),
    ("health_insurance", "Krankenversicherung (sensibel)", False),
    ("guardian1_first_name", "Erziehungsberechtigter 1 – Vorname", True),
    ("guardian1_last_name", "Erziehungsberechtigter 1 – Nachname", True),
    ("guardian1_phone", "Erziehungsberechtigter 1 – Telefon", False),
    ("guardian1_email", "Erziehungsberechtigter 1 – E-Mail", False),
    ("guardian2_first_name", "Erziehungsberechtigter 2 – Vorname", False),
    ("guardian2_last_name", "Erziehungsberechtigter 2 – Nachname", False),
    ("guardian2_phone", "Erziehungsberechtigter 2 – Telefon", False),
    ("guardian2_email", "Erziehungsberechtigter 2 – E-Mail", False),
]

COLUMN_ALIASES = {
    "sequential_number": ["lfd nr", "lfd. nr.", "laufende nummer", "nummer", "mitgliedsnummer"],
    "first_name": ["vorname", "mitglied vorname"],
    "last_name": ["nachname", "name", "familienname", "mitglied nachname"],
    "gender": ["geschlecht"],
    "entry_date": ["beitrittsdatum", "eintrittsdatum", "eintritt", "beitritt"],
    "jf_card_number": ["jf ausweisnummer", "jf-ausweisnummer", "ausweisnummer"],
    "activity_status": ["aktivität", "aktivitaet", "status"],
    "birth_date": ["geburtsdatum", "geburtstag", "geb.-datum", "geb datum", "geburts datum"],
    "street": ["straße", "strasse", "straße und hausnummer", "anschrift", "adresse"],
    "postal_code": ["plz", "postleitzahl"],
    "city": ["ort", "wohnort"],
    "phone": ["telefon", "festnetz"],
    "mobile": ["mobil", "handy", "mobiltelefon"],
    "email": ["e-mail", "email", "e-mail-adresse", "mail"],
    "photo_video_permission": ["foto/video", "fotoeinwilligung", "foto-/videoeinwilligung", "foto erlaubt"],
    "divera_permission": ["divera", "divera einwilligung", "divera-einwilligung"],
    "school_or_employer": ["schule", "arbeitgeber", "schule oder arbeitgeber", "schul-/berufsverhältnis"],
    "vocational_training": ["berufsausbildung", "ausbildung", "aktuelle ausbildung"],
    "career_goal": ["berufsziel", "voraussichtliches berufsziel"],
    "other_organizations": ["weitere vereine", "andere vereine", "organisationen", "vereine und organisationen"],
    "health_notes": ["krankheiten", "allergien", "gesundheitliche hinweise", "krankheiten behinderungen beschwerden allergien"],
    "swimming_status": ["schwimmer", "schwimmfähigkeit", "schwimmer/nichtschwimmer"],
    "health_insurance": ["krankenversicherung", "krankenkasse"],
    "guardian1_first_name": ["erziehungsberechtigter 1 vorname", "kontakt 1 vorname", "eb1 vorname", "mutter vorname"],
    "guardian1_last_name": ["erziehungsberechtigter 1 nachname", "kontakt 1 nachname", "eb1 nachname", "mutter nachname"],
    "guardian1_phone": ["erziehungsberechtigter 1 telefon", "kontakt 1 telefon", "eb1 telefon", "mutter telefon"],
    "guardian1_email": ["erziehungsberechtigter 1 e-mail", "kontakt 1 e-mail", "eb1 email", "mutter email"],
    "guardian2_first_name": ["erziehungsberechtigter 2 vorname", "kontakt 2 vorname", "eb2 vorname", "vater vorname"],
    "guardian2_last_name": ["erziehungsberechtigter 2 nachname", "kontakt 2 nachname", "eb2 nachname", "vater nachname"],
    "guardian2_phone": ["erziehungsberechtigter 2 telefon", "kontakt 2 telefon", "eb2 telefon", "vater telefon"],
    "guardian2_email": ["erziehungsberechtigter 2 e-mail", "kontakt 2 e-mail", "eb2 email", "vater email"],
}

STATUS_MAP = {
    "aktiv": Member.ActivityStatus.ACTIVE,
    "active": Member.ActivityStatus.ACTIVE,
    "beurlaubt": Member.ActivityStatus.ON_LEAVE,
    "on leave": Member.ActivityStatus.ON_LEAVE,
    "übergang": Member.ActivityStatus.TRANSITION,
    "uebergang": Member.ActivityStatus.TRANSITION,
    "transition": Member.ActivityStatus.TRANSITION,
    "austritt": Member.ActivityStatus.RESIGNED,
    "ausgetreten": Member.ActivityStatus.RESIGNED,
    "resigned": Member.ActivityStatus.RESIGNED,
}
GENDER_MAP = {
    "m": Member.Gender.MALE,
    "männlich": Member.Gender.MALE,
    "maennlich": Member.Gender.MALE,
    "w": Member.Gender.FEMALE,
    "weiblich": Member.Gender.FEMALE,
    "d": Member.Gender.DIVERSE,
    "divers": Member.Gender.DIVERSE,
    "keine angabe": Member.Gender.NOT_SPECIFIED,
}

SWIMMING_MAP = {
    "schwimmer": Member.SwimmingStatus.SWIMMER,
    "ja": Member.SwimmingStatus.SWIMMER,
    "nichtschwimmer": Member.SwimmingStatus.NON_SWIMMER,
    "nein": Member.SwimmingStatus.NON_SWIMMER,
    "keine angabe": Member.SwimmingStatus.NOT_SPECIFIED,
}


def normalize_header(value: Any) -> str:
    value = str(value or "").strip().lower().replace("_", " ")
    value = value.replace("\n", " ").replace("\r", " ")
    value = re.sub(r"[.:;]+$", "", value)
    value = re.sub(r"\s+", " ", value)
    return value


def detect_columns(headers: list[Any]) -> dict[str, int]:
    normalized = [normalize_header(value) for value in headers]
    mapping: dict[str, int] = {}
    for field, aliases in COLUMN_ALIASES.items():
        candidates = {normalize_header(field), *(normalize_header(alias) for alias in aliases)}
        for index, header in enumerate(normalized):
            if header in candidates:
                mapping[field] = index
                break
    return mapping


def inspect_workbook(path: str | Path) -> dict[str, Any]:
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"Die Excel-Datei konnte nicht gelesen werden: {exc}") from exc
    sheets = []
    for sheet in workbook.worksheets:
        sample_rows = []
        for row in sheet.iter_rows(min_row=1, max_row=10, values_only=True):
            sample_rows.append([clean_text(value) for value in row[:20]])
        sheets.append({
            "name": sheet.title,
            "max_row": sheet.max_row or len(sample_rows),
            "max_column": sheet.max_column or max((len(row) for row in sample_rows), default=0),
            "sample_rows": sample_rows,
        })
    return {"sheets": sheets}


def read_headers(path: str | Path, sheet_name: str, header_row: int) -> tuple[list[str], list[list[str]]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError("Das gewählte Tabellenblatt existiert nicht mehr.")
    sheet = workbook[sheet_name]
    if header_row < 1:
        raise ValueError("Die gewählte Kopfzeile ist ungültig.")
    # In read-only mode, openpyxl may return None for max_row when worksheet
    # dimensions are missing or were written by another spreadsheet program.
    # Reading the requested row directly is therefore more reliable than
    # comparing against sheet.max_row.
    row = next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
    headers = []
    for index, value in enumerate(row):
        label = clean_text(value)
        headers.append(label or f"Spalte {get_column_letter(index + 1)}")
    while headers and not clean_text(row[len(headers) - 1]):
        headers.pop()
    if not headers:
        raise ValueError("In der gewählten Kopfzeile wurden keine Spaltenbezeichnungen gefunden.")
    preview_rows = []
    for values in sheet.iter_rows(min_row=header_row + 1, max_row=header_row + 5, values_only=True):
        preview_rows.append([clean_text(value) for value in values[: len(headers)]])
    return headers, preview_rows


def parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Ungültiges Datum: {text}")


def parse_bool(value: Any) -> bool:
    return normalize_header(value) in {"ja", "j", "yes", "true", "1", "x", "erteilt"}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def next_free_numbers(count: int) -> list[int]:
    used = set(Member.objects.values_list("sequential_number", flat=True))
    result = []
    candidate = 1
    while len(result) < count:
        if candidate not in used:
            result.append(candidate)
            used.add(candidate)
        candidate += 1
    return result


def parse_member_workbook(path_or_upload, *, sheet_name: str | None = None, header_row: int = 1, mapping: dict[str, int] | None = None) -> dict[str, Any]:
    try:
        workbook = load_workbook(path_or_upload, data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"Die Excel-Datei konnte nicht gelesen werden: {exc}") from exc
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    rows = list(sheet.iter_rows(min_row=header_row, values_only=True))
    if not rows:
        raise ValueError("Die Excel-Datei enthält ab der gewählten Kopfzeile keine Daten.")
    headers = list(rows[0])
    mapping = mapping if mapping is not None else detect_columns(headers)
    missing = [field for field in ("first_name", "last_name", "birth_date", "guardian1_first_name", "guardian1_last_name") if field not in mapping]
    if missing:
        label_map = dict((name, label) for name, label, _required in IMPORT_FIELDS)
        raise ValueError("Folgende Pflichtfelder wurden nicht zugeordnet: " + ", ".join(label_map[x] for x in missing))
    used_columns = [index for index in mapping.values()]
    if len(used_columns) != len(set(used_columns)):
        raise ValueError("Eine Excel-Spalte wurde mehreren JF-Manager-Feldern zugeordnet. Bitte jede Spalte nur einmal verwenden.")

    raw_items = []
    for excel_row, row in enumerate(rows[1:], start=header_row + 1):
        if not any(value not in (None, "") for value in row):
            continue
        item = {field: (row[index] if index < len(row) else None) for field, index in mapping.items()}
        item["excel_row"] = excel_row
        raw_items.append(item)

    auto_numbers = iter(next_free_numbers(sum(1 for item in raw_items if not item.get("sequential_number"))))
    preview = []
    seen_numbers: set[int] = set()
    seen_cards: set[str] = set()
    seen_people: set[tuple[str, str, str]] = set()
    for item in raw_items:
        errors: list[str] = []
        warnings: list[str] = []
        try:
            birth_date = parse_date(item.get("birth_date"))
        except ValueError as exc:
            birth_date = None
            errors.append(str(exc))
        try:
            entry_date = parse_date(item.get("entry_date"))
        except ValueError as exc:
            entry_date = None
            errors.append(str(exc))
        first_name = clean_text(item.get("first_name"))
        last_name = clean_text(item.get("last_name"))
        if not first_name:
            errors.append("Vorname fehlt")
        if not last_name:
            errors.append("Nachname fehlt")
        if not birth_date:
            errors.append("Geburtsdatum fehlt oder ist ungültig")

        number_raw = clean_text(item.get("sequential_number"))
        if number_raw:
            try:
                sequential_number = int(float(number_raw.replace(",", ".")))
                if sequential_number <= 0:
                    raise ValueError
            except ValueError:
                sequential_number = None
                errors.append("Laufende Nummer ist ungültig")
        else:
            sequential_number = next(auto_numbers)
            warnings.append(f"Laufende Nummer automatisch vergeben: {sequential_number}")

        g1_first = clean_text(item.get("guardian1_first_name"))
        g1_last = clean_text(item.get("guardian1_last_name"))
        g1_phone = clean_text(item.get("guardian1_phone"))
        g1_email = clean_text(item.get("guardian1_email"))
        if not g1_first or not g1_last:
            errors.append("Erste erziehungsberechtigte Person fehlt")
        elif not g1_phone and not g1_email:
            errors.append("Kontaktmöglichkeit der ersten erziehungsberechtigten Person fehlt")

        duplicate = None
        if first_name and last_name and birth_date:
            duplicate = Member.objects.filter(first_name__iexact=first_name, last_name__iexact=last_name, birth_date=birth_date).first()
        jf_number = clean_text(item.get("jf_card_number"))
        if not duplicate and jf_number:
            duplicate = Member.objects.filter(jf_card_number=jf_number).first()
        if not duplicate and sequential_number:
            duplicate = Member.objects.filter(sequential_number=sequential_number).first()
        if duplicate:
            warnings.append(f"Mögliches Duplikat: {duplicate}")

        person_key = (first_name.casefold(), last_name.casefold(), birth_date.isoformat() if birth_date else "")
        if sequential_number in seen_numbers:
            errors.append(f"Laufende Nummer {sequential_number} kommt in der Datei mehrfach vor")
        if jf_number and jf_number.casefold() in seen_cards:
            errors.append(f"JF-Ausweisnummer {jf_number} kommt in der Datei mehrfach vor")
        if birth_date and person_key in seen_people:
            errors.append("Mitglied mit gleichem Namen und Geburtsdatum kommt in der Datei mehrfach vor")
        if sequential_number:
            seen_numbers.add(sequential_number)
        if jf_number:
            seen_cards.add(jf_number.casefold())
        if birth_date:
            seen_people.add(person_key)

        status_text = normalize_header(item.get("activity_status"))
        gender_text = normalize_header(item.get("gender"))
        data = {
            "excel_row": item["excel_row"],
            "sequential_number": sequential_number,
            "first_name": first_name,
            "last_name": last_name,
            "gender": GENDER_MAP.get(gender_text, Member.Gender.NOT_SPECIFIED if gender_text else ""),
            "entry_date": entry_date.isoformat() if entry_date else None,
            "jf_card_number": jf_number or None,
            "activity_status": STATUS_MAP.get(status_text, Member.ActivityStatus.ACTIVE),
            "birth_date": birth_date.isoformat() if birth_date else None,
            "street": clean_text(item.get("street")),
            "postal_code": clean_text(item.get("postal_code")),
            "city": clean_text(item.get("city")),
            "phone": clean_text(item.get("phone")),
            "mobile": clean_text(item.get("mobile")),
            "email": clean_text(item.get("email")),
            "photo_video_permission": parse_bool(item.get("photo_video_permission")),
            "divera_permission": parse_bool(item.get("divera_permission")),
            "school_or_employer": clean_text(item.get("school_or_employer")),
            "vocational_training": clean_text(item.get("vocational_training")),
            "career_goal": clean_text(item.get("career_goal")),
            "other_organizations": clean_text(item.get("other_organizations")),
            "health_notes": clean_text(item.get("health_notes")),
            "swimming_status": SWIMMING_MAP.get(normalize_header(item.get("swimming_status")), Member.SwimmingStatus.NOT_SPECIFIED),
            "health_insurance": clean_text(item.get("health_insurance")),
            "guardians": [{"priority": 1, "first_name": g1_first, "last_name": g1_last, "phone": g1_phone, "email": g1_email}],
        }
        g2_first = clean_text(item.get("guardian2_first_name"))
        g2_last = clean_text(item.get("guardian2_last_name"))
        g2_phone = clean_text(item.get("guardian2_phone"))
        g2_email = clean_text(item.get("guardian2_email"))
        if any((g2_first, g2_last, g2_phone, g2_email)):
            if not g2_first or not g2_last:
                errors.append("Zweite erziehungsberechtigte Person ist unvollständig")
            elif not g2_phone and not g2_email:
                errors.append("Kontaktmöglichkeit der zweiten erziehungsberechtigten Person fehlt")
            data["guardians"].append({"priority": 2, "first_name": g2_first, "last_name": g2_last, "phone": g2_phone, "email": g2_email})
        preview.append({"data": data, "selected": True, "errors": errors, "warnings": warnings, "duplicate_id": duplicate.pk if duplicate else None})
    return {"mapping": mapping, "headers": [clean_text(x) for x in headers], "rows": preview, "total": len(preview)}



def validate_edited_preview_rows(post, original_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Rebuild and validate editable preview rows from submitted form data."""
    rows: list[dict[str, Any]] = []
    seen_numbers: set[int] = set()
    seen_cards: set[str] = set()
    seen_people: set[tuple[str, str, str]] = set()

    for index, original in enumerate(original_rows):
        prefix = f"row-{index}-"
        selected = post.get(prefix + "selected") == "on"
        data = dict(original.get("data", {}))
        data["excel_row"] = original.get("data", {}).get("excel_row")

        text_fields = (
            "first_name", "last_name", "jf_card_number", "street", "postal_code", "city",
            "phone", "mobile", "email", "school_or_employer", "vocational_training",
            "career_goal", "other_organizations", "health_notes", "health_insurance",
        )
        for field in text_fields:
            data[field] = clean_text(post.get(prefix + field, ""))

        data["gender"] = post.get(prefix + "gender", "")
        data["activity_status"] = post.get(prefix + "activity_status", Member.ActivityStatus.ACTIVE)
        data["photo_video_permission"] = post.get(prefix + "photo_video_permission") == "on"
        data["divera_permission"] = post.get(prefix + "divera_permission") == "on"
        swimming = post.get(prefix + "swimming_status", Member.SwimmingStatus.NOT_SPECIFIED)
        data["swimming_status"] = swimming if swimming in Member.SwimmingStatus.values else Member.SwimmingStatus.NOT_SPECIFIED

        errors: list[str] = []
        warnings: list[str] = []
        number_raw = clean_text(post.get(prefix + "sequential_number", ""))
        try:
            data["sequential_number"] = int(number_raw) if number_raw else None
            if data["sequential_number"] is not None and data["sequential_number"] <= 0:
                raise ValueError
        except ValueError:
            data["sequential_number"] = None
            errors.append("Laufende Nummer ist ungültig")

        for field, label in (("birth_date", "Geburtsdatum"), ("entry_date", "Beitrittsdatum")):
            raw = clean_text(post.get(prefix + field, ""))
            try:
                parsed = parse_date(raw)
                data[field] = parsed.isoformat() if parsed else None
            except ValueError:
                data[field] = None
                errors.append(f"{label} ist ungültig")

        guardians = []
        for priority in (1, 2):
            guardian = {
                "priority": priority,
                "first_name": clean_text(post.get(prefix + f"guardian{priority}_first_name", "")),
                "last_name": clean_text(post.get(prefix + f"guardian{priority}_last_name", "")),
                "phone": clean_text(post.get(prefix + f"guardian{priority}_phone", "")),
                "email": clean_text(post.get(prefix + f"guardian{priority}_email", "")),
            }
            if priority == 1 or any(guardian[key] for key in ("first_name", "last_name", "phone", "email")):
                guardians.append(guardian)
        data["guardians"] = guardians

        if selected:
            if not data["first_name"]:
                errors.append("Vorname fehlt")
            if not data["last_name"]:
                errors.append("Nachname fehlt")
            if not data["birth_date"]:
                errors.append("Geburtsdatum fehlt oder ist ungültig")
            if not data["sequential_number"]:
                errors.append("Laufende Nummer fehlt oder ist ungültig")

            first_guardian = guardians[0] if guardians else None
            if not first_guardian or not first_guardian["first_name"] or not first_guardian["last_name"]:
                errors.append("Erste erziehungsberechtigte Person fehlt")
            elif not first_guardian["phone"] and not first_guardian["email"]:
                errors.append("Kontaktmöglichkeit der ersten erziehungsberechtigten Person fehlt")
            if len(guardians) > 1:
                second = guardians[1]
                if not second["first_name"] or not second["last_name"]:
                    errors.append("Zweite erziehungsberechtigte Person ist unvollständig")
                elif not second["phone"] and not second["email"]:
                    errors.append("Kontaktmöglichkeit der zweiten erziehungsberechtigten Person fehlt")

        duplicate = None
        birth = parse_date(data["birth_date"]) if data.get("birth_date") else None
        if selected and data["first_name"] and data["last_name"] and birth:
            duplicate = Member.objects.filter(first_name__iexact=data["first_name"], last_name__iexact=data["last_name"], birth_date=birth).first()
        if selected and not duplicate and data.get("jf_card_number"):
            duplicate = Member.objects.filter(jf_card_number=data["jf_card_number"]).first()
        if selected and not duplicate and data.get("sequential_number"):
            duplicate = Member.objects.filter(sequential_number=data["sequential_number"]).first()
        if duplicate:
            warnings.append(f"Mögliches Duplikat: {duplicate}")

        if selected:
            number = data.get("sequential_number")
            card = (data.get("jf_card_number") or "").casefold()
            person = (data["first_name"].casefold(), data["last_name"].casefold(), data.get("birth_date") or "")
            if number in seen_numbers:
                errors.append(f"Laufende Nummer {number} kommt in der Auswahl mehrfach vor")
            if card and card in seen_cards:
                errors.append(f"JF-Ausweisnummer {data['jf_card_number']} kommt in der Auswahl mehrfach vor")
            if data.get("birth_date") and person in seen_people:
                errors.append("Mitglied mit gleichem Namen und Geburtsdatum kommt in der Auswahl mehrfach vor")
            if number:
                seen_numbers.add(number)
            if card:
                seen_cards.add(card)
            if data.get("birth_date"):
                seen_people.add(person)

        rows.append({
            "data": data,
            "selected": selected,
            "errors": errors if selected else [],
            "warnings": warnings if selected else [],
            "duplicate_id": duplicate.pk if duplicate and selected else None,
        })
    return rows

def import_preview_rows(rows: list[dict[str, Any]]) -> dict[str, int]:
    created = skipped = failed = 0
    for row in rows:
        if not row.get("selected", True) or row.get("errors") or row.get("duplicate_id"):
            skipped += 1
            continue
        data = row["data"].copy()
        guardians = data.pop("guardians")
        data.pop("excel_row", None)
        try:
            with transaction.atomic():
                member = Member.objects.create(**data)
                for guardian in guardians:
                    Guardian.objects.create(member=member, **guardian)
            created += 1
        except Exception:
            failed += 1
    return {"created": created, "skipped": skipped, "failed": failed}


def create_import_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mitgliederimport"
    headers = [
        "Lfd. Nr.", "Vorname", "Nachname", "Geschlecht", "Geburtsdatum", "Beitrittsdatum",
        "JF-Ausweisnummer", "Aktivität", "Straße und Hausnummer", "PLZ", "Ort", "Telefon", "Mobil",
        "E-Mail", "Foto-/Videoeinwilligung", "DIVERA-Einwilligung",
        "Schule oder Arbeitgeber", "Berufsausbildung", "Voraussichtliches Berufsziel",
        "Weitere Vereine oder Organisationen", "Gesundheitliche Hinweise", "Schwimmfähigkeit", "Krankenversicherung",
        "Erziehungsberechtigter 1 Vorname", "Erziehungsberechtigter 1 Nachname",
        "Erziehungsberechtigter 1 Telefon", "Erziehungsberechtigter 1 E-Mail",
        "Erziehungsberechtigter 2 Vorname", "Erziehungsberechtigter 2 Nachname",
        "Erziehungsberechtigter 2 Telefon", "Erziehungsberechtigter 2 E-Mail",
    ]
    ws.append(headers)
    ws.append(["", "Max", "Mustermann", "männlich", date(2012, 5, 14), date(2024, 1, 10), "JF-001", "Aktiv", "Musterstraße 1", "12345", "Musterstadt", "", "0170 1234567", "max@example.de", "Ja", "Ja", "IGS Musterstadt", "", "Feuerwehrmann", "Sportverein", "Haselnussallergie", "Schwimmer", "Musterkasse", "Erika", "Mustermann", "06305 1234", "erika@example.de", "", "", "", ""])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="B42335")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 42
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20
    for cell in (ws["E2"], ws["F2"]):
        cell.number_format = "DD.MM.YYYY"
    notes = wb.create_sheet("Hinweise")
    notes.append(["Hinweis", "Beschreibung"])
    notes.append(["Pflichtfelder", "Vorname, Nachname, Geburtsdatum sowie mindestens eine erziehungsberechtigte Person."])
    notes.append(["Laufende Nummer", "Kann leer bleiben; der JF-Manager vergibt dann automatisch die nächste freie Nummer."])
    notes.append(["Datumsformat", "TT.MM.JJJJ oder echtes Excel-Datum."])
    notes.append(["Aktivität", "Aktiv, Beurlaubt, Übergang oder Austritt."])
    notes.append(["Einwilligungen", "Ja, Nein, X, 1 oder Erteilt."])
    notes.append(["Freiwillige Angaben", "Schule/Beruf, weitere Organisationen sowie Gesundheits- und Versicherungsangaben sind freiwillig."])
    notes.append(["Gesundheitsdaten", "Besonders schützenswert. Nur importieren, wenn eine passende Rechtsgrundlage bzw. ausdrückliche Einwilligung vorliegt."])
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
