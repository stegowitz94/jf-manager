from datetime import date
from html import escape
from io import BytesIO

from django.http import FileResponse, HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def age_on(birth_date, reference=None):
    if not birth_date:
        return ""
    reference = reference or date.today()
    return reference.year - birth_date.year - ((reference.month, reference.day) < (birth_date.month, birth_date.day))


def guardian_values(member, priority):
    guardian = next((item for item in member.guardians.all() if item.priority == priority), None)
    if not guardian:
        return ("", "", "")
    return (f"{guardian.first_name} {guardian.last_name}", guardian.phone, guardian.email)


def export_members_xlsx(members, label="Mitgliederliste"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Mitglieder"
    headers = [
        "Lfd. Nr.", "Vorname", "Nachname", "Geschlecht", "Geburtsdatum", "Alter",
        "Beitrittsdatum", "JF-Ausweisnummer", "Aktivität", "Straße und Hausnummer", "PLZ", "Ort", "Telefon",
        "Mobil", "E-Mail", "Foto/Video erlaubt", "DIVERA erlaubt",
        "Schule oder Arbeitgeber", "Berufsausbildung", "Voraussichtliches Berufsziel", "Weitere Vereine/Organisationen",
        "Erziehungsberechtigte Person 1", "Telefon 1", "E-Mail 1",
        "Erziehungsberechtigte Person 2", "Telefon 2", "E-Mail 2",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="B42335")
        cell.alignment = Alignment(vertical="center")
    for member in members:
        g1 = guardian_values(member, 1)
        g2 = guardian_values(member, 2)
        ws.append([
            member.sequential_number, member.first_name, member.last_name, member.get_gender_display(),
            member.birth_date, age_on(member.birth_date), member.entry_date, member.jf_card_number or "",
            member.get_activity_status_display(), member.street, member.postal_code, member.city, member.phone, member.mobile,
            member.email, "Ja" if member.photo_video_permission else "Nein",
            "Ja" if member.divera_permission else "Nein", member.school_or_employer,
            member.vocational_training, member.career_goal, member.other_organizations, *g1, *g2,
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28
    date_columns = {5, 7}
    for row in ws.iter_rows(min_row=2):
        for index in date_columns:
            row[index - 1].number_format = "DD.MM.YYYY"
    widths = [11, 18, 20, 14, 14, 8, 15, 18, 14, 25, 9, 18, 18, 18, 26, 18, 15, 28, 24, 28, 32, 28, 18, 28, 28, 18, 28]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"JF-Mitgliederliste_{date.today().isoformat()}.xlsx"
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _pdf_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawString(15 * mm, 10 * mm, f"JF-Manager · Seite {doc.page}")
    canvas.restoreState()


def export_members_pdf(members, export_type="compact"):
    stream = BytesIO()
    doc = SimpleDocTemplate(
        stream, pagesize=landscape(A4), rightMargin=10 * mm, leftMargin=10 * mm,
        topMargin=12 * mm, bottomMargin=16 * mm, title="JF-Mitgliederliste",
    )
    styles = getSampleStyleSheet()
    story = [Paragraph("Jugendfeuerwehr – Mitgliederliste", styles["Title"]), Spacer(1, 5 * mm)]
    if export_type == "contacts":
        headers = ["Mitglied", "Telefon", "Mobil", "E-Mail", "Erziehungsberechtigte", "Kontakt"]
        rows = [headers]
        for member in members:
            guardian_lines = []
            contact_lines = []
            for guardian in member.guardians.all():
                guardian_lines.append(f"{guardian.first_name} {guardian.last_name}")
                contact_lines.append(" / ".join(filter(None, [guardian.phone, guardian.email])))
            rows.append([
                f"{member.first_name} {member.last_name}", member.phone, member.mobile, member.email,
                "\n".join(guardian_lines), "\n".join(contact_lines),
            ])
        widths = [40, 29, 29, 46, 53, 70]
    elif export_type == "full":
        headers = ["Nr.", "Name", "Geburtsdatum", "Status", "Anschrift", "Kontakt", "Einwilligungen"]
        rows = [headers]
        for member in members:
            rows.append([
                member.sequential_number,
                f"{member.first_name} {member.last_name}",
                member.birth_date.strftime("%d.%m.%Y"), member.get_activity_status_display(),
                f"{member.street}\n{member.postal_code} {member.city}".strip(),
                " / ".join(filter(None, [member.phone, member.mobile, member.email])),
                f"Foto/Video: {'Ja' if member.photo_video_permission else 'Nein'}\nDIVERA: {'Ja' if member.divera_permission else 'Nein'}",
            ])
        widths = [13, 42, 28, 27, 46, 76, 41]
    else:
        headers = ["Nr.", "Vorname", "Nachname", "Geburtsdatum", "Alter", "Status", "Mobil", "Ort"]
        rows = [headers]
        for member in members:
            rows.append([
                member.sequential_number, member.first_name, member.last_name,
                member.birth_date.strftime("%d.%m.%Y"), age_on(member.birth_date),
                member.get_activity_status_display(), member.mobile or member.phone, member.city,
            ])
        widths = [14, 37, 42, 28, 14, 28, 37, 43]
    # Paragraphs allow wrapping and line breaks.
    wrapped = [[Paragraph(escape(str(value or "")).replace("\n", "<br/>"), styles["BodyText"]) for value in row] for row in rows]
    table = Table(wrapped, colWidths=[width * mm for width in widths], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#B42335")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7BDC5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F8")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    stream.seek(0)
    return FileResponse(stream, as_attachment=True, filename=f"JF-Mitgliederliste_{date.today().isoformat()}.pdf")
