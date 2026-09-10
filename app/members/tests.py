from datetime import date

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse

from .models import Guardian, Leave, Member, MemberStatusHistory


class MemberModelTests(TestCase):
    def setUp(self):
        self.member = Member.objects.create(
            sequential_number=1,
            first_name="Max",
            last_name="Mustermann",
            birth_date=date(2012, 5, 3),
        )

    def test_initial_status_history_is_created(self):
        history = MemberStatusHistory.objects.get(member=self.member)
        self.assertEqual(history.status, Member.ActivityStatus.ACTIVE)

    def test_status_change_creates_history_entry(self):
        self.member.activity_status = Member.ActivityStatus.ON_LEAVE
        self.member.save()
        self.assertEqual(MemberStatusHistory.objects.filter(member=self.member).count(), 2)

    def test_guardian_requires_contact_method(self):
        guardian = Guardian(member=self.member, priority=1, first_name="Erika", last_name="Mustermann")
        with self.assertRaises(ValidationError):
            guardian.full_clean()

    def test_leave_end_must_not_precede_start(self):
        leave = Leave(member=self.member, start_date=date(2026, 8, 10), planned_end_date=date(2026, 8, 1))
        with self.assertRaises(ValidationError):
            leave.full_clean()


class MemberViewTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="wart", password="testpass123")
        self.client.login(username="wart", password="testpass123")

    def test_member_list_requires_and_accepts_login(self):
        response = self.client.get(reverse("members:list"))
        self.assertEqual(response.status_code, 200)

    def test_member_create_with_guardian(self):
        response = self.client.post(
            reverse("members:create"),
            {
                "sequential_number": 2,
                "first_name": "Anna",
                "last_name": "Beispiel",
                "birth_date": "2013-02-01",
                "activity_status": Member.ActivityStatus.ACTIVE,
                "guardians-TOTAL_FORMS": "2",
                "guardians-INITIAL_FORMS": "0",
                "guardians-MIN_NUM_FORMS": "1",
                "guardians-MAX_NUM_FORMS": "2",
                "guardians-0-priority": "1",
                "guardians-0-first_name": "Eva",
                "guardians-0-last_name": "Beispiel",
                "guardians-0-phone": "01234 5678",
                "guardians-0-email": "",
                "guardians-1-priority": "",
                "guardians-1-first_name": "",
                "guardians-1-last_name": "",
                "guardians-1-phone": "",
                "guardians-1-email": "",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Member.objects.filter(first_name="Anna").exists())

class MemberExportTests(TestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model
        self.user = get_user_model().objects.create_user(username="export", password="secret-test-password")
        self.client.force_login(self.user)
        self.member = Member.objects.create(
            sequential_number=999,
            first_name="Export",
            last_name="Test",
            birth_date="2012-01-01",
        )
        Guardian.objects.create(
            member=self.member, priority=1, first_name="Eltern", last_name="Test", phone="01234 5678"
        )

    def test_excel_export(self):
        response = self.client.post(reverse("members:export"), {"format": "xlsx", "statuses": ["active"]})
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])

    def test_pdf_export(self):
        response = self.client.post(
            reverse("members:export"), {"format": "pdf", "template": "compact", "statuses": ["active"]}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")

class MemberEnhancementTests(TestCase):
    def test_member_can_store_street(self):
        member = Member.objects.create(
            sequential_number=991,
            first_name="Straßen",
            last_name="Test",
            birth_date=date(2012, 1, 1),
            street="Musterstraße 1",
        )
        self.assertEqual(member.street, "Musterstraße 1")

class MemberImportMappingTests(TestCase):
    def test_detect_columns_recognizes_geburtstag(self):
        from .imports import detect_columns
        mapping = detect_columns(["Vorname", "Nachname", "Geburtstag", "Mutter Vorname", "Mutter Nachname"])
        self.assertEqual(mapping["birth_date"], 2)
        self.assertEqual(mapping["guardian1_first_name"], 3)

    def test_manual_mapping_can_read_nonstandard_headers(self):
        from tempfile import NamedTemporaryFile
        from openpyxl import Workbook
        from .imports import parse_member_workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Jugend"
        sheet.append(["Rufname", "Familie", "Geb.", "Kontakt A", "Kontakt B", "Kontakt Tel"])
        sheet.append(["Anna", "Beispiel", "12.03.2012", "Erika", "Beispiel", "0170 1234567"])
        with NamedTemporaryFile(suffix=".xlsx") as temp:
            workbook.save(temp.name)
            result = parse_member_workbook(
                temp.name,
                sheet_name="Jugend",
                header_row=1,
                mapping={
                    "first_name": 0,
                    "last_name": 1,
                    "birth_date": 2,
                    "guardian1_first_name": 3,
                    "guardian1_last_name": 4,
                    "guardian1_phone": 5,
                },
            )
        self.assertEqual(result["total"], 1)
        self.assertFalse(result["rows"][0]["errors"])
        self.assertEqual(result["rows"][0]["data"]["first_name"], "Anna")


class MemberOptionalSensitiveDataTests(TestCase):
    def test_optional_school_and_health_fields_can_be_saved(self):
        member = Member.objects.create(
            sequential_number=992,
            first_name="Freiwillig",
            last_name="Test",
            birth_date=date(2012, 2, 2),
            school_or_employer="Musterschule",
            career_goal="Notfallsanitäter",
            other_organizations="Sportverein",
            health_notes="Pollenallergie",
            swimming_status=Member.SwimmingStatus.SWIMMER,
            health_insurance="Musterkasse",
        )
        self.assertEqual(member.school_or_employer, "Musterschule")
        self.assertEqual(member.swimming_status, Member.SwimmingStatus.SWIMMER)

    def test_standard_excel_export_does_not_include_health_data(self):
        from openpyxl import load_workbook
        from io import BytesIO
        user = get_user_model().objects.create_user(username="privacy", password="secret-test-password")
        self.client.force_login(user)
        Member.objects.create(
            sequential_number=993,
            first_name="Privat",
            last_name="Test",
            birth_date=date(2012, 3, 3),
            health_notes="Darf nicht exportiert werden",
            health_insurance="Vertraulich",
        )
        response = self.client.post(reverse("members:export"), {"format": "xlsx", "statuses": ["active"]})
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        headers = [cell.value for cell in next(workbook.active.iter_rows(min_row=1, max_row=1))]
        self.assertNotIn("Gesundheitliche Hinweise", headers)
        self.assertNotIn("Krankenversicherung", headers)
